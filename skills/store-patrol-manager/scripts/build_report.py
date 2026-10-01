#!/usr/bin/env python3
"""店铺勾选字段环比周报生成脚本。

用法：
    python3 build_report.py \
      --current <本期query结果JSON> \
      --prev <上期query结果JSON> \
      --output <输出xlsx路径> \
      [--label 本期标签] [--prev-label 上期标签]

两个 JSON 文件内容为 query_shop_metrics 返回的完整文本
（含 schema_version 等外层字段均可，脚本自动截取首个 '{' 起解析）。

输出 Excel 三个 Sheet：
  - 周期汇总：指标 | 本期值 | 上期值 | 汇总环比
  - 每日明细：指标 | D1 | D2 | ...（本期各天）
  - 每日环比：指标 | D2环比 | D3环比 | ...（较前一天）
"""
import argparse
import json
import re
import sys
from pathlib import Path
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少依赖 openpyxl，请先执行: pip3 install openpyxl")

# workspace 根 = 本脚本上级三级（skills/<skill>/scripts -> workspace 根）。
# 相对路径统一基于它解析，保证从任何 cwd / 任何设备运行输出位置一致。
WS_ROOT = Path(__file__).resolve().parents[3]


def resolve(p):
    """相对路径基于 workspace 根解析（不依赖 cwd）；绝对路径原样返回。"""
    if not p:
        return p
    p = Path(p)
    return str(p if p.is_absolute() else WS_ROOT / p)

# 用户勾选的 17 个字段（key -> 中文名）
SELECTED_FIELDS = [
    ("payAmt", "支付金额"),
    ("uv", "访客数"),
    ("payByrCnt", "支付买家数"),
    ("payRate", "支付转化率"),
    ("pv", "浏览量"),
    ("payItmCnt", "支付件数"),
    ("payShopRfdAmt", "店铺退款金额"),
    ("netPaymentAmount", "净支付金额"),
    ("payOrdCnt", "支付订单数"),
    ("payPct", "客单价"),
    ("cartByrCnt", "加购人数"),
    ("cartCnt", "加购件数"),
    ("p4pExpendAmt", "关键词推广花费"),
    ("cubeAmt", "精准人群花费"),
    ("admCostFamtQzt", "全站推花费"),
    ("olderPayAmt", "老客支付金额"),
    ("payOldByrCnt", "老客支付买家数"),
]

PERCENT_KEYS = {"payRate"}
MONEY_KEYS = {
    "payAmt", "netPaymentAmount", "payShopRfdAmt", "payPct",
    "p4pExpendAmt", "cubeAmt", "admCostFamtQzt", "olderPayAmt",
}


def load_result(path):
    """读取 query_shop_metrics 返回文本，解析 JSON，返回 {period, daily}."""
    with open(path, encoding="utf-8") as f:
        content = f.read()
    start = content.find("{")
    if start < 0:
        sys.exit(f"无法在 {path} 中找到 JSON")
    data = json.loads(content[start:])
    period = {}
    daily = []
    for table in data.get("tables", []):
        view = table.get("output_view")
        if view == "period_summary" and table.get("rows"):
            period = table["rows"][0].get("data", {})
        elif view == "daily_trend":
            for row in table.get("rows", []):
                cells = {c["key"]: c for c in row.get("cells", [])}
                daily.append((row["day_key"], cells))
    return period, daily


def kind_of(key):
    if key in PERCENT_KEYS:
        return "percent"
    if key in MONEY_KEYS:
        return "money"
    return "number"


def fmt_val(v, kind):
    if v is None:
        return "-"
    try:
        v = float(v)
    except (TypeError, ValueError):
        return v
    if kind == "percent":
        return f"{v*100:.2f}%"
    if kind == "money":
        return f"{v:,.2f}"
    if v == int(v):
        return f"{int(v):,}"
    return f"{v:,.2f}"


def fmt_chg(cur, prev):
    """环比百分比字符串：上涨带 +，下跌带 -；无基准返回 -."""
    if cur is None or prev is None:
        return "-"
    try:
        cur = float(cur)
        prev = float(prev)
    except (TypeError, ValueError):
        return "-"
    if prev == 0:
        return "新增" if cur != 0 else "-"
    return f"{(cur - prev) / abs(prev) * 100:+.2f}%"


def calc_chg_float(cur, prev):
    """返回环比数值（小数），无基准返回 None."""
    if cur is None or prev is None:
        return None
    try:
        cur = float(cur)
        prev = float(prev)
    except (TypeError, ValueError):
        return None
    if prev == 0:
        return None
    return (cur - prev) / abs(prev)


# ---------- 样式 ----------
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
UP_FONT = Font(name="Arial", size=10, color="C00000")
DOWN_FONT = Font(name="Arial", size=10, color="008000")
TITLE_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FILL = PatternFill("solid", start_color="2E75B6")
ALT_FILL = PatternFill("solid", start_color="F2F7FC")
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def style_chg_cell(cell, chg_str):
    cell.border = BORDER
    cell.alignment = RIGHT
    if chg_str not in ("-", "新增") and chg_str:
        cell.font = UP_FONT if chg_str.startswith("+") else DOWN_FONT
    else:
        cell.font = BODY_FONT


def build(current_path, prev_path, output_path, label, prev_label):
    cur_period, cur_daily = load_result(current_path)
    prev_period, _ = load_result(prev_path) if prev_path else ({}, [])
    days = [d for d, _ in cur_daily]

    wb = Workbook()

    # ===== Sheet1 周期汇总 =====
    ws = wb.active
    ws.title = "周期汇总"
    ws.merge_cells("A1:D1")
    ws["A1"] = f"店铺勾选字段 · 周汇总（{label}）"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 26

    headers = ["指标", label, prev_label, "汇总环比"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[2].height = 20

    ri = 3
    for key, lname in SELECTED_FIELDS:
        cv = cur_period.get(key, {}).get("value") if isinstance(cur_period.get(key), dict) else cur_period.get(key)
        pv = prev_period.get(key, {}).get("value") if isinstance(prev_period.get(key), dict) else prev_period.get(key)
        kind = kind_of(key)
        name_cell = ws.cell(row=ri, column=1, value=lname)
        cur_cell = ws.cell(row=ri, column=2, value=fmt_val(cv, kind))
        prev_cell = ws.cell(row=ri, column=3, value=fmt_val(pv, kind))
        chg_cell = ws.cell(row=ri, column=4, value=fmt_chg(cv, pv))
        for c in (name_cell, cur_cell, prev_cell, chg_cell):
            c.border = BORDER
            if ri % 2 == 1:
                c.fill = ALT_FILL
        name_cell.font = BODY_FONT
        name_cell.alignment = LEFT
        cur_cell.font = BODY_FONT
        cur_cell.alignment = RIGHT
        prev_cell.font = BODY_FONT
        prev_cell.alignment = RIGHT
        style_chg_cell(chg_cell, chg_cell.value)
        ri += 1

    note = ws.cell(row=ri + 1, column=1,
                   value=f"注：汇总环比 = (本期值 - {prev_label}) / {prev_label} × 100%；上涨标红，下跌标绿；上期无数据时为 -。")
    note.font = Font(name="Arial", size=9, italic=True, color="808080")

    ws.column_dimensions["A"].width = 22
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    # ===== Sheet2 每日明细 =====
    ws2 = wb.create_sheet("每日明细")
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(days))
    ws2.cell(row=1, column=1, value=f"每日明细（{label}）").font = TITLE_FONT
    ws2.cell(row=1, column=1).fill = TITLE_FILL
    ws2.cell(row=1, column=1).alignment = CENTER
    ws2.row_dimensions[1].height = 26

    hdr = ["指标"] + [d[5:] + "日" for d in days]
    for ci, h in enumerate(hdr, 1):
        c = ws2.cell(row=2, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws2.row_dimensions[2].height = 20

    ri = 3
    for key, lname in SELECTED_FIELDS:
        name_cell = ws2.cell(row=ri, column=1, value=lname)
        name_cell.font = BODY_FONT
        name_cell.border = BORDER
        name_cell.alignment = LEFT
        if ri % 2 == 1:
            name_cell.fill = ALT_FILL
        for di, (day, cells) in enumerate(cur_daily):
            raw = cells.get(key)
            v = raw.get("value") if isinstance(raw, dict) else raw
            val_cell = ws2.cell(row=ri, column=2 + di, value=fmt_val(v, kind_of(key)))
            val_cell.font = BODY_FONT
            val_cell.border = BORDER
            val_cell.alignment = RIGHT
            if ri % 2 == 1:
                val_cell.fill = ALT_FILL
        ri += 1
    ws2.column_dimensions["A"].width = 22
    for di in range(len(days)):
        ws2.column_dimensions[get_column_letter(2 + di)].width = 13

    # ===== Sheet3 每日环比 =====
    ws3 = wb.create_sheet("每日环比")
    ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + max(len(days) - 1, 1))
    ws3.cell(row=1, column=1, value="每日环比（较前一天，%）：上涨标红，下跌标绿").font = TITLE_FONT
    ws3.cell(row=1, column=1).fill = TITLE_FILL
    ws3.cell(row=1, column=1).alignment = CENTER
    ws3.row_dimensions[1].height = 26

    hdr3 = ["指标"] + [d[5:] + "日环比" for d in days[1:]]
    for ci, h in enumerate(hdr3, 1):
        c = ws3.cell(row=2, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws3.row_dimensions[2].height = 20

    ri = 3
    for key, lname in SELECTED_FIELDS:
        name_cell = ws3.cell(row=ri, column=1, value=lname)
        name_cell.font = BODY_FONT
        name_cell.border = BORDER
        name_cell.alignment = LEFT
        if ri % 2 == 1:
            name_cell.fill = ALT_FILL
        for di in range(1, len(days)):
            prev_raw = cur_daily[di - 1][1].get(key)
            cur_raw = cur_daily[di][1].get(key)
            pv = prev_raw.get("value") if isinstance(prev_raw, dict) else prev_raw
            cv = cur_raw.get("value") if isinstance(cur_raw, dict) else cur_raw
            chg_str = fmt_chg(cv, pv)
            val_cell = ws3.cell(row=ri, column=1 + di, value=chg_str)
            style_chg_cell(val_cell, chg_str)
            if ri % 2 == 1:
                val_cell.fill = ALT_FILL
        ri += 1
    ws3.column_dimensions["A"].width = 22
    for di in range(1, len(days)):
        ws3.column_dimensions[get_column_letter(1 + di)].width = 14

    ws.freeze_panes = "A3"
    ws2.freeze_panes = "B3"
    ws3.freeze_panes = "B3"

    wb.save(output_path)
    print(f"saved: {output_path}")
    print(f"days: {len(days)} ({days[0]} ~ {days[-1]})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--current", required=True, help="本期 query_shop_metrics 返回文本文件")
    ap.add_argument("--prev", default="", help="上期 query_shop_metrics 返回文本文件（可选）")
    ap.add_argument("--output", required=True, help="输出 xlsx 路径")
    ap.add_argument("--label", default="本期", help="本期标签，如 2026-08-24~08-30")
    ap.add_argument("--prev-label", default="上期", help="上期标签")
    args = ap.parse_args()
    build(resolve(args.current), resolve(args.prev) or None, resolve(args.output), args.label, args.prev_label)
